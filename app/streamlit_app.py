import logging
import sys
from pathlib import Path

import streamlit as st


# ==========================================
# 1. CONFIGURACION DEL PROYECTO
# ==========================================

ROOT = Path(__file__).resolve().parents[1]

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.dashboard import mostrar_dashboard
from app.file_utils import preparar_descargas
from src.agent.gemini_agent import preguntar_agente

from src.database.sales_queries import (
    obtener_pedidos,
    obtener_detalle_productos
)

from src.reports.excel_generator import (
    generar_reporte_excel
)

from openpyxl import load_workbook

from src.reports.personalized_report import (
    generar_reporte_personalizado as crear_excel_personalizado
)

CARPETA_REPORTES = (ROOT / "reports").resolve()

MIME_EXCEL = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


# ==========================================
# 2. CONFIGURACION DE LA PAGINA
# ==========================================

st.set_page_config(
    page_title="DataInsight AI",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ==========================================
# 3. ESTADO DE LA APLICACION
# ==========================================

if "mensajes" not in st.session_state:
    st.session_state.mensajes = []

if "reporte_base" not in st.session_state:
    st.session_state.reporte_base = None

if "reporte_personalizado" not in st.session_state:
    st.session_state.reporte_personalizado = None


# ==========================================
# 4. FUNCIONES AUXILIARES
# ==========================================

def mostrar_mensaje(mensaje, indice):
    """Muestra un mensaje y sus descargas."""

    with st.chat_message(mensaje["role"]):

        st.markdown(mensaje["content"])

        for numero, archivo in enumerate(
            mensaje.get("archivos", [])
        ):

            st.download_button(
                label="📥 Descargar reporte Excel",
                data=archivo["contenido"],
                file_name=archivo["nombre"],
                mime=MIME_EXCEL,
                key=f"descarga_{indice}_{numero}"
            )


# ==========================================
# 5. PANEL LATERAL
# ==========================================

with st.sidebar:

    st.title("📊 DataInsight AI")

    st.caption(
        "Plataforma de análisis empresarial "
        "asistida por inteligencia artificial."
    )

    st.divider()

    # Selector de seccion
    seccion = st.radio(
        "Navegación",
        options=[
            "📊 Dashboard",
            "💬 Asistente IA"
        ],
        key="seccion_activa"
    )

    st.divider()

    # Generador independiente de Gemini
    st.subheader("Reportes Excel")

    st.caption(
        "Genera un reporte directamente desde "
        "SQL Server, sin utilizar Gemini."
    )

    if st.button(
        "📊 Generar Excel base (sin IA)",
        key="generar_excel_base",
        use_container_width=True
    ):

        try:

            with st.spinner("Generando reporte..."):

                pedidos = obtener_pedidos()
                detalle = obtener_detalle_productos()

                ruta = generar_reporte_excel(
                    pedidos=pedidos,
                    detalle=detalle
                )

                if not ruta.is_file():
                    raise RuntimeError(
                        "No se encontro el Excel generado."
                    )

                st.session_state.reporte_base = {
                    "nombre": ruta.name,
                    "contenido": ruta.read_bytes()
                }

            st.success(
                "Reporte generado correctamente."
            )

        except Exception:

            logging.exception(
                "Error generando el reporte base"
            )

            st.error(
                "No se pudo generar el reporte. "
                "Revisa la terminal."
            )

    # Descargar ultimo reporte base
    reporte_base = st.session_state.reporte_base

    if reporte_base is not None:

        st.download_button(
            label="📥 Descargar Excel base",
            data=reporte_base["contenido"],
            file_name=reporte_base["nombre"],
            mime=MIME_EXCEL,
            key="descargar_excel_base",
            use_container_width=True
        )

        st.caption(
            "Último reporte base generado correctamente. "
            "Incluye 5 hojas y 2 gráficos, sin IA."
        )


    # ======================================
    # GENERADOR DE EXCEL PERSONALIZADO
    # ======================================

    st.divider()

    with st.expander(
        "🧩 Excel personalizado",
        expanded=True
    ):

        st.caption(
            "Selecciona un año, un territorio "
            "o ambos para generar un reporte."
        )

        # 1. Formulario de filtros
        with st.form("form_reporte_personalizado"):

            anio_seleccionado = st.selectbox(
                "Año",
                options=[
                    "Todos",
                    "2011",
                    "2012",
                    "2013",
                    "2014"
                ],
                index=0
            )

            territorio_ingresado = st.text_input(
                "Territorio (opcional)",
                placeholder="Ejemplo: Canada"
            )

            incluir_ia = st.checkbox(
                "Incluir conclusiones de IA",
                value=False
            )

            enviar = st.form_submit_button(
                "📊 Generar Excel personalizado",
                use_container_width=True
            )

        # 2. Procesar formulario
        if enviar:

            # Evitar mostrar una descarga anterior
            # como resultado de la nueva solicitud
            st.session_state.reporte_personalizado = None

            anio = (
                None
                if anio_seleccionado == "Todos"
                else int(anio_seleccionado)
            )

            territorio = (
                territorio_ingresado.strip() or None
            )

            # Exigir al menos un filtro
            if anio is None and territorio is None:

                st.warning(
                    "Selecciona un año o escribe "
                    "un territorio para continuar."
                )

            else:

                try:

                    with st.spinner(
                        "Generando Excel personalizado..."
                    ):

                        ruta = crear_excel_personalizado(
                            anio=anio,
                            territorio=territorio,
                            incluir_conclusiones_ia=incluir_ia
                        )

                        if not ruta.is_file():
                            raise RuntimeError(
                                "El archivo no fue generado."
                            )

                        # Verificar si realmente se
                        # incorporaron conclusiones IA
                        wb = load_workbook(
                            ruta,
                            read_only=True
                        )

                        try:
                            tiene_ia = (
                                "Conclusiones IA"
                                in wb.sheetnames
                            )
                        finally:
                            wb.close()

                        # Guardar archivo para descarga
                        st.session_state.reporte_personalizado = {
                            "nombre": ruta.name,
                            "contenido": ruta.read_bytes(),
                            "tiene_ia": tiene_ia
                        }

                    st.success(
                        "Excel personalizado generado "
                        "correctamente."
                    )

                    if incluir_ia and not tiene_ia:

                        st.warning(
                            "Se generó el Excel con las "
                            "cinco hojas analíticas, pero "
                            "no pudieron incorporarse "
                            "las conclusiones de Gemini."
                        )

                except ValueError as error:

                    st.warning(
                        f"No se pudo generar el reporte: "
                        f"{error}"
                    )

                except Exception:

                    logging.exception(
                        "Error generando Excel personalizado"
                    )

                    st.error(
                        "Ocurrió un error al generar "
                        "el reporte personalizado. "
                        "Revisa la terminal."
                    )

        # 3. Mostrar ultima descarga disponible
        reporte = st.session_state.get(
            "reporte_personalizado"
        )

        if reporte is not None:

            st.download_button(
                label="📥 Descargar Excel personalizado",
                data=reporte["contenido"],
                file_name=reporte["nombre"],
                mime=MIME_EXCEL,
                key="descargar_excel_personalizado",
                use_container_width=True
            )

            if reporte["tiene_ia"]:

                st.caption(
                    "El archivo incluye seis hojas, "
                    "con conclusiones de Gemini."
                )

            else:

                st.caption(
                    "El archivo incluye cinco hojas "
                    "analíticas, sin conclusiones IA."
                )


    st.divider()

    if seccion == "💬 Asistente IA":

        st.subheader("Conversación")

        if st.button(
            "🗑️ Limpiar conversación",
            use_container_width=True
        ):

            st.session_state.mensajes = []
            st.rerun()

        st.caption(
            "Ejemplos: productos más vendidos, "
            "territorios líderes o generación "
            "de reportes Excel."
        )

    st.divider()

    st.caption(
        "Python · SQL Server · Pandas · "
        "Gemini · Streamlit"
    )


# ==========================================
# 6. ENCABEZADO PRINCIPAL
# ==========================================

st.title("DataInsight AI")

st.caption(
    "Análisis de ventas de AdventureWorks2022 "
    "con indicadores, visualizaciones y Gemini."
)

st.divider()


# ==========================================
# 7. SECCION DASHBOARD
# ==========================================

if seccion == "📊 Dashboard":

    st.subheader("Panel de indicadores")

    st.write(
        "Consulta los resultados de ventas "
        "y explora su evolución."
    )

    mostrar_dashboard()


# ==========================================
# 8. SECCION ASISTENTE IA
# ==========================================

elif seccion == "💬 Asistente IA":

    st.subheader("Asistente inteligente")

    st.caption(
        "Realiza preguntas sobre las ventas "
        "o solicita reportes Excel."
    )

    # Mostrar historial
    if not st.session_state.mensajes:

        st.info(
            "Puedes comenzar preguntando por "
            "indicadores, productos o territorios."
        )

    for indice, mensaje in enumerate(
        st.session_state.mensajes
    ):
        mostrar_mensaje(mensaje, indice)

    # Entrada del chat
    pregunta = st.chat_input(
        "Escribe una pregunta sobre las ventas..."
    )

    if pregunta:

        mensaje_usuario = {
            "role": "user",
            "content": pregunta
        }

        st.session_state.mensajes.append(
            mensaje_usuario
        )

        with st.chat_message("user"):
            st.markdown(pregunta)

        with st.chat_message("assistant"):

            reportes_generados = []
            ocurrio_error = False

            with st.spinner(
                "Gemini está analizando tu solicitud..."
            ):

                try:

                    respuesta = preguntar_agente(
                        pregunta,
                        on_reporte=reportes_generados.append
                    )

                except Exception:

                    logging.exception(
                        "Error procesando la solicitud"
                    )

                    ocurrio_error = True

                    if reportes_generados:

                        respuesta = (
                            "El reporte Excel fue generado, "
                            "pero no se pudo completar la "
                            "respuesta final de Gemini. "
                            "Puedes descargar el archivo "
                            "a continuación."
                        )

                    else:

                        respuesta = (
                            "No se pudo completar la "
                            "solicitud. El servicio de IA "
                            "podría estar temporalmente "
                            "no disponible."
                        )

            # Preparar descargas asociadas a
            # esta pregunta
            archivos = preparar_descargas(
                reportes_generados,
                CARPETA_REPORTES
            )

            mensaje_agente = {
                "role": "assistant",
                "content": respuesta,
                "archivos": archivos
            }

            st.session_state.mensajes.append(
                mensaje_agente
            )

            if ocurrio_error:
                st.warning(respuesta)
            else:
                st.markdown(respuesta)

            # Mostrar botones de descarga
            indice = (
                len(st.session_state.mensajes) - 1
            )

            for numero, archivo in enumerate(
                archivos
            ):

                st.download_button(
                    label="📥 Descargar reporte Excel",
                    data=archivo["contenido"],
                    file_name=archivo["nombre"],
                    mime=MIME_EXCEL,
                    key=f"descarga_{indice}_{numero}"
                )