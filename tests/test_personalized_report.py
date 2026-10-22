import pandas as pd
import pytest

from openpyxl import load_workbook

from src.reports import excel_generator
from src.reports import personalized_report


# ==========================================
# DATOS SIMULADOS
# ==========================================

@pytest.fixture
def entorno_simulado(monkeypatch, tmp_path):

    pedidos = pd.DataFrame({
        "SalesOrderID": [101, 102, 103, 104],
        "OrderDate": [
            "2012-01-10",
            "2013-02-15",
            "2013-06-20",
            "2014-03-10"
        ],
        "SubTotal": [100, 200, 300, 400],
        "Territorio": [
            "Southwest",
            "Southwest",
            "Canada",
            "Canada"
        ]
    })

    detalle = pd.DataFrame({
        "SalesOrderID": [
            101, 101, 102, 103, 103, 104
        ],
        "ProductID": [1, 2, 3, 4, 5, 6],
        "NombreProducto": [
            "Producto A",
            "Producto B",
            "Producto C",
            "Producto D",
            "Producto E",
            "Producto F"
        ],
        "OrderQty": [1, 1, 1, 1, 1, 1],
        "LineTotal": [
            60, 40, 200, 120, 180, 400
        ]
    })

    # Evitar consultas reales a SQL Server
    monkeypatch.setattr(
        personalized_report,
        "obtener_pedidos",
        lambda: pedidos.copy()
    )

    monkeypatch.setattr(
        personalized_report,
        "obtener_detalle_productos",
        lambda: detalle.copy()
    )

    # Guardar Excel en carpeta temporal
    monkeypatch.setattr(
        excel_generator,
        "ROOT",
        tmp_path
    )


# ==========================================
# PRUEBA 1: REPORTE GENERAL
# ==========================================

def test_generar_reporte_general(
    entorno_simulado
):

    ruta = (
        personalized_report
        .generar_reporte_personalizado()
    )

    assert ruta.is_file()

    wb = load_workbook(ruta)

    assert wb["Resumen"]["B5"].value == 4

    assert wb["Resumen"]["B6"].value == (
        pytest.approx(1000)
    )

    assert len(wb.sheetnames) == 5

    wb.close()


# ==========================================
# PRUEBA 2: FILTRO POR AÑO
# ==========================================

def test_generar_reporte_2013(
    entorno_simulado
):

    ruta = (
        personalized_report
        .generar_reporte_personalizado(
            anio=2013
        )
    )

    assert ruta.is_file()
    assert "_2013.xlsx" in ruta.name

    wb = load_workbook(ruta)

    assert wb["Resumen"]["B5"].value == 2

    assert wb["Resumen"]["B6"].value == (
        pytest.approx(500)
    )

    graficos = sum(
        len(hoja._charts)
        for hoja in wb.worksheets
    )

    assert graficos == 2

    wb.close()


# ==========================================
# PRUEBA 3: FILTRO POR TERRITORIO
# ==========================================

def test_generar_reporte_canada(
    entorno_simulado
):

    ruta = (
        personalized_report
        .generar_reporte_personalizado(
            territorio="Canada"
        )
    )

    assert ruta.is_file()
    assert "_Canada.xlsx" in ruta.name

    wb = load_workbook(ruta)

    assert wb["Resumen"]["B5"].value == 2

    assert wb["Resumen"]["B6"].value == (
        pytest.approx(700)
    )

    wb.close()


# ==========================================
# PRUEBA 4: DOS FILTROS
# ==========================================

def test_generar_reporte_2013_canada(
    entorno_simulado
):

    ruta = (
        personalized_report
        .generar_reporte_personalizado(
            anio=2013,
            territorio="Canada"
        )
    )

    assert ruta.is_file()
    assert "_2013_Canada.xlsx" in ruta.name

    wb = load_workbook(ruta)

    assert wb["Resumen"]["B5"].value == 1

    assert wb["Resumen"]["B6"].value == (
        pytest.approx(300)
    )

    territorios = wb["Territorios"]

    assert territorios["A2"].value == "Canada"

    assert territorios["C2"].value == (
        pytest.approx(300)
    )

    wb.close()


# ==========================================
# PRUEBA 5: SIN RESULTADOS
# ==========================================

def test_no_generar_excel_sin_datos(
    entorno_simulado
):

    with pytest.raises(
        ValueError,
        match="No se encontraron pedidos"
    ):

        personalized_report.generar_reporte_personalizado(
            anio=2020
        )


# ==========================================
# PRUEBA 6: CONCLUSIONES CON DATOS FILTRADOS
# ==========================================

def test_conclusiones_ia_respetan_filtros(
    entorno_simulado,
    monkeypatch
):

    datos_recibidos = []

    # Simular Gemini sin consumir la API
    def simular_gemini(pedidos, detalle):

        datos_recibidos.append({
            "ids_pedidos": set(
                pedidos["SalesOrderID"]
            ),
            "ids_detalle": set(
                detalle["SalesOrderID"]
            ),
            "ventas": float(
                pedidos["SubTotal"].sum()
            )
        })

        return [
            "Primera conclusion filtrada.",
            "Segunda conclusion filtrada.",
            "Tercera conclusion filtrada."
        ]

    monkeypatch.setattr(
        personalized_report,
        "generar_conclusiones_ia",
        simular_gemini
    )

    # Generar reporte Canada 2013 con IA
    ruta = (
        personalized_report
        .generar_reporte_personalizado(
            anio=2013,
            territorio="Canada",
            incluir_conclusiones_ia=True
        )
    )

    assert ruta.is_file()

    # Verificar lo que recibio Gemini
    assert len(datos_recibidos) == 1

    assert datos_recibidos[0][
        "ids_pedidos"
    ] == {103}

    assert datos_recibidos[0][
        "ids_detalle"
    ] == {103}

    assert datos_recibidos[0][
        "ventas"
    ] == pytest.approx(300)

    # Verificar las seis hojas
    wb = load_workbook(ruta)

    assert len(wb.sheetnames) == 6

    assert "Conclusiones IA" in wb.sheetnames

    assert wb["Resumen"]["B6"].value == (
        pytest.approx(300)
    )

    hoja = wb["Conclusiones IA"]

    assert (
        "Primera conclusion filtrada."
        in hoja["A4"].value
    )

    wb.close()


# ==========================================
# PRUEBA 7: FALLA GEMINI Y SE CONSERVA EXCEL
# ==========================================

def test_reporte_filtrado_si_falla_gemini(
    entorno_simulado,
    monkeypatch
):

    def simular_falla(pedidos, detalle):
        raise RuntimeError(
            "Error simulado de Gemini"
        )

    monkeypatch.setattr(
        personalized_report,
        "generar_conclusiones_ia",
        simular_falla
    )

    ruta = (
        personalized_report
        .generar_reporte_personalizado(
            anio=2013,
            territorio="Canada",
            incluir_conclusiones_ia=True
        )
    )

    # El archivo debe existir aunque Gemini falle
    assert ruta.is_file()

    wb = load_workbook(ruta)

    assert len(wb.sheetnames) == 5

    assert "Conclusiones IA" not in wb.sheetnames

    # Los datos filtrados deben mantenerse
    assert wb["Resumen"]["B5"].value == 1

    assert wb["Resumen"]["B6"].value == (
        pytest.approx(300)
    )

    wb.close()