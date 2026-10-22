from contextlib import contextmanager
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from google.genai import types

from src.agent import gemini_agent, tools


# ==========================================
# GENERADOR FICTICIO PARA LAS PRUEBAS
# ==========================================

@pytest.fixture
def generador_simulado(monkeypatch, tmp_path):

    llamadas = []
    opciones = {"crear_hoja_ia": True}

    def simular_generacion(
        anio=None,
        territorio=None,
        incluir_conclusiones_ia=False
    ):

        llamadas.append({
            "anio": anio,
            "territorio": territorio,
            "incluir_conclusiones_ia":
                incluir_conclusiones_ia
        })

        ruta = tmp_path / (
            f"reporte_prueba_{len(llamadas)}.xlsx"
        )

        wb = Workbook()
        wb.active.title = "Resumen"

        for nombre in [
            "Ventas mensuales",
            "Territorios",
            "Top ingresos",
            "Top unidades"
        ]:
            wb.create_sheet(nombre)

        if (
            incluir_conclusiones_ia
            and opciones["crear_hoja_ia"]
        ):
            wb.create_sheet("Conclusiones IA")

        wb.save(ruta)
        wb.close()

        return ruta

    monkeypatch.setattr(
        tools.personalized_report,
        "generar_reporte_personalizado",
        simular_generacion
    )

    return llamadas, opciones


# ==========================================
# PRUEBA 1: REPORTE PERSONALIZADO CON IA
# ==========================================

def test_herramienta_reporte_con_ia(
    generador_simulado
):

    llamadas, _ = generador_simulado

    resultado = gemini_agent.ejecutar_herramienta(
        "generar_reporte_personalizado",
        {
            "anio": 2013,
            "territorio": "Canada",
            "incluir_conclusiones_ia": True
        }
    )

    assert resultado["estado"] == "generado"
    assert resultado["tipo_reporte"] == "personalizado"

    assert resultado["filtros"] == {
        "anio": 2013,
        "territorio": "Canada"
    }

    assert resultado["cantidad_hojas"] == 6
    assert resultado["conclusiones_ia_en_excel"] is True

    assert len(llamadas) == 1

    assert llamadas[0] == {
        "anio": 2013,
        "territorio": "Canada",
        "incluir_conclusiones_ia": True
    }


# ==========================================
# PRUEBA 2: REPORTE SOLO POR AÑO
# ==========================================

def test_herramienta_reporte_sin_ia(
    generador_simulado
):

    llamadas, _ = generador_simulado

    resultado = gemini_agent.ejecutar_herramienta(
        "generar_reporte_personalizado",
        {"anio": 2013}
    )

    assert resultado["cantidad_hojas"] == 5
    assert resultado["conclusiones_ia_en_excel"] is False

    assert llamadas[0]["anio"] == 2013
    assert llamadas[0]["territorio"] is None


# ==========================================
# PRUEBA 3: RECHAZAR SOLICITUD SIN FILTROS
# ==========================================

def test_herramienta_exige_filtros(
    generador_simulado
):

    with pytest.raises(ValueError):

        gemini_agent.ejecutar_herramienta(
            "generar_reporte_personalizado",
            {}
        )


# ==========================================
# PRUEBA 4: RECHAZAR ARGUMENTOS INVALIDOS
# ==========================================

def test_herramienta_rechaza_argumentos(
    generador_simulado
):

    casos_invalidos = [
        {"anio": "2013"},
        {"anio": True},
        {"territorio": " "},
        {
            "anio": 2013,
            "incluir_conclusiones_ia": "si"
        },
        {
            "anio": 2013,
            "argumento_no_autorizado": 123
        }
    ]

    for argumentos in casos_invalidos:

        with pytest.raises(ValueError):

            gemini_agent.ejecutar_herramienta(
                "generar_reporte_personalizado",
                argumentos
            )


# ==========================================
# PRUEBA 5: IA FALLA, EXCEL CONSERVADO
# ==========================================

def test_herramienta_advierte_si_falla_ia(
    generador_simulado
):

    _, opciones = generador_simulado

    # Simular que no se incorpora la sexta hoja
    opciones["crear_hoja_ia"] = False

    resultado = gemini_agent.ejecutar_herramienta(
        "generar_reporte_personalizado",
        {
            "anio": 2013,
            "incluir_conclusiones_ia": True
        }
    )

    assert resultado["estado"] == "generado"
    assert resultado["cantidad_hojas"] == 5

    assert resultado["conclusiones_ia_en_excel"] is False
    assert "advertencia" in resultado


# ==========================================
# PRUEBA 6: CACHE Y AVISO A STREAMLIT
# ==========================================

def test_agente_no_duplica_reporte(
    monkeypatch
):

    argumentos = {
        "anio": 2013,
        "territorio": "Canada",
        "incluir_conclusiones_ia": True
    }

    nombre = "generar_reporte_personalizado"

    llamada = SimpleNamespace(
        name=nombre,
        args=argumentos
    )

    ejecuciones = []
    avisos = []
    solicitudes_modelo = []

    reporte_ficticio = {
        "estado": "generado",
        "archivo": "reporte_prueba.xlsx",
        "carpeta": "reports",
        "conclusiones_ia_en_excel": True
    }

    # Simular la ejecucion de la herramienta
    def ejecutar_simulado(
        nombre_herramienta,
        argumentos_herramienta
    ):

        ejecuciones.append({
            "nombre": nombre_herramienta,
            "argumentos": argumentos_herramienta
        })

        return reporte_ficticio

    monkeypatch.setattr(
        gemini_agent,
        "ejecutar_herramienta",
        ejecutar_simulado
    )

    # Simular respuestas de Gemini
    def respuesta_simulada(**kwargs):

        solicitudes_modelo.append(kwargs)

        if len(solicitudes_modelo) == 1:

            contenido = types.Content(
                role="model",
                parts=[
                    types.Part.from_function_call(
                        name=nombre,
                        args=argumentos
                    ),
                    types.Part.from_function_call(
                        name=nombre,
                        args=argumentos
                    )
                ]
            )

            return SimpleNamespace(
                function_calls=[llamada, llamada],
                candidates=[
                    SimpleNamespace(content=contenido)
                ],
                text=None
            )

        return SimpleNamespace(
            function_calls=[],
            candidates=[],
            text="Reporte personalizado generado."
        )

    @contextmanager
    def cliente_simulado():

        yield SimpleNamespace(
            models=SimpleNamespace(
                generate_content=respuesta_simulada
            )
        )

    monkeypatch.setattr(
        gemini_agent,
        "crear_cliente_gemini",
        cliente_simulado
    )

    monkeypatch.setattr(
        gemini_agent,
        "obtener_modelo",
        lambda: "modelo_simulado"
    )

    # Ejecutar agente sin conectarse a Gemini
    respuesta = gemini_agent.preguntar_agente(
        "Genera un Excel de Canada de 2013.",
        on_reporte=avisos.append
    )

    assert respuesta == (
        "Reporte personalizado generado."
    )

    # Aunque la solicitud se repita,
    # generar el Excel una sola vez
    assert len(ejecuciones) == 1

    # Notificar a Streamlit solo una vez
    assert len(avisos) == 1

    assert avisos[0] == reporte_ficticio