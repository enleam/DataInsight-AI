import pandas as pd
import pytest

from openpyxl import load_workbook

from src.reports import excel_generator


# DATOS FICTICIOS PARA LAS PRUEBAS

@pytest.fixture
def datos_ventas():

    pedidos = pd.DataFrame({
        "SalesOrderID": [1, 2, 3],
        "OrderDate": [
            "2023-01-05",
            "2023-02-06",
            "2023-02-07"
        ],
        "SubTotal": [100, 200, 300],
        "Territorio": [
            "Norte",
            "Sur",
            "Norte"
        ]
    })

    detalle = pd.DataFrame({
        "SalesOrderID": [1, 2, 3],
        "ProductID": [10, 20, 10],
        "NombreProducto": [
            "Producto A",
            "Producto B",
            "Producto A"
        ],
        "OrderQty": [2, 1, 3],
        "LineTotal": [100, 200, 300]
    })

    return pedidos, detalle


# PRUEBA 1: ESTRUCTURA DEL EXCEL

def test_reporte_cinco_hojas_y_dos_graficos(
    tmp_path, monkeypatch, datos_ventas
):

    monkeypatch.setattr(
        excel_generator,
        "ROOT",
        tmp_path
    )

    pedidos, detalle = datos_ventas

    ruta = excel_generator.generar_reporte_excel(
        pedidos,
        detalle
    )

    assert ruta.exists()

    wb = load_workbook(ruta)

    assert wb.sheetnames == [
        "Resumen",
        "Ventas mensuales",
        "Territorios",
        "Top ingresos",
        "Top unidades"
    ]

    cantidad_graficos = sum(
        len(hoja._charts)
        for hoja in wb.worksheets
    )

    assert cantidad_graficos == 2

    wb.close()


# PRUEBA 2: INDICADORES Y DATOS

def test_indicadores_del_excel(
    tmp_path, monkeypatch, datos_ventas
):

    monkeypatch.setattr(
        excel_generator,
        "ROOT",
        tmp_path
    )

    pedidos, detalle = datos_ventas

    ruta = excel_generator.generar_reporte_excel(
        pedidos,
        detalle
    )

    wb = load_workbook(ruta)

    resumen = wb["Resumen"]

    assert resumen["B5"].value == 3
    assert resumen["B6"].value == pytest.approx(600)
    assert resumen["B7"].value == pytest.approx(200)

    territorios = wb["Territorios"]

    assert territorios["A2"].value == "Norte"
    assert territorios["C2"].value == pytest.approx(400)

    productos = wb["Top ingresos"]

    assert productos["B2"].value == "Producto A"
    assert productos["E2"].value == pytest.approx(400)

    wb.close()


# PRUEBA 3: CONCLUSIONES DE IA

def test_agregar_conclusiones_ia(
    tmp_path, monkeypatch, datos_ventas
):

    monkeypatch.setattr(
        excel_generator,
        "ROOT",
        tmp_path
    )

    pedidos, detalle = datos_ventas

    ruta = excel_generator.generar_reporte_excel(
        pedidos,
        detalle
    )

    conclusiones = [
        "Las ventas acumuladas fueron 600.",
        "El territorio Norte lidera los ingresos.",
        "El Producto A genera mayores ingresos."
    ]

    excel_generator.agregar_conclusiones_ia(
        ruta,
        conclusiones
    )

    wb = load_workbook(ruta)

    assert len(wb.sheetnames) == 6

    assert "Conclusiones IA" in wb.sheetnames

    hoja = wb["Conclusiones IA"]

    for i, conclusion in enumerate(
        conclusiones,
        start=4
    ):
        assert conclusion in hoja[f"A{i}"].value

    cantidad_graficos = sum(
        len(hoja._charts)
        for hoja in wb.worksheets
    )

    assert cantidad_graficos == 2

    wb.close()


# PRUEBA 4: CONCLUSIONES INVALIDAS

def test_rechazar_conclusiones_invalidas(
    tmp_path, monkeypatch, datos_ventas
):

    monkeypatch.setattr(
        excel_generator,
        "ROOT",
        tmp_path
    )

    pedidos, detalle = datos_ventas

    ruta = excel_generator.generar_reporte_excel(
        pedidos,
        detalle
    )

    with pytest.raises(ValueError):
        excel_generator.agregar_conclusiones_ia(
            ruta,
            ["Solo una conclusion"]
        )

    wb = load_workbook(ruta)

    assert len(wb.sheetnames) == 5

    wb.close()


# PRUEBA 5: CONSISTENCIA DE VENTAS

def test_rechazar_ventas_inconsistentes(
    tmp_path, monkeypatch, datos_ventas
):

    monkeypatch.setattr(
        excel_generator,
        "ROOT",
        tmp_path
    )

    pedidos, detalle = datos_ventas

    detalle = detalle.copy()

    # Crear una diferencia intencional
    detalle.loc[0, "LineTotal"] = 150

    with pytest.raises(
        ValueError,
        match="no coinciden"
    ):
        excel_generator.generar_reporte_excel(
            pedidos,
            detalle
        )