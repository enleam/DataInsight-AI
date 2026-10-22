from openpyxl import load_workbook

from src.reports import personalized_report

from src.reports.excel_generator import (
    generar_reporte_excel,
    agregar_conclusiones_ia
)

from src.agent.insights import generar_conclusiones_ia

from src.database.sales_queries import (
    obtener_pedidos,
    obtener_detalle_productos
)

from src.analytics.sales_analysis import (
    preparar_pedidos,
    calcular_kpis,
    ventas_por_territorio,
    analizar_productos
)


def consultar_indicadores_generales() -> dict:
    """Obtiene los KPIs generales de ventas."""

    df = preparar_pedidos(obtener_pedidos())

    if df.empty:
        return {"mensaje": "No existen pedidos."}

    kpis = calcular_kpis(df)

    return {
        "total_pedidos": int(kpis["total_pedidos"]),
        "ventas_productos": round(
            float(kpis["total_ventas"]), 2
        ),
        "ticket_promedio": round(
            float(kpis["ticket_promedio"]), 2
        ),
        "primera_fecha": str(
            kpis["primera_fecha"].date()
        ),
        "ultima_fecha": str(
            kpis["ultima_fecha"].date()
        ),
        "metrica": "SubTotal",
        "moneda": "no especificada"
    }


def consultar_ventas_por_territorio(limite: int) -> dict:
    """Obtiene los territorios con mayores ingresos."""

    if type(limite) is not int or not 1 <= limite <= 10:
        raise ValueError(
            "El limite debe ser un entero entre 1 y 10."
        )

    df = preparar_pedidos(obtener_pedidos())

    if df.empty:
        return {"territorios": [], "mensaje": "Sin datos"}

    resumen = ventas_por_territorio(df)
    top = resumen.head(limite)

    resultados = []

    for _, fila in top.iterrows():
        resultados.append({
            "territorio": str(fila["Territorio"]),
            "cantidad_pedidos": int(
                fila["CantidadPedidos"]
            ),
            "ventas_productos": round(
                float(fila["VentasProductos"]), 2
            )
        })

    return {
        "metrica": "SubTotal",
        "moneda": "no especificada",
        "periodo_inicio": str(
            df["OrderDate"].min().date()
        ),
        "periodo_fin": str(
            df["OrderDate"].max().date()
        ),
        "territorios": resultados
    }


def consultar_top_productos(
    limite: int,
    criterio: str
) -> dict:
    """Obtiene los productos lideres por ingresos o unidades."""

    if type(limite) is not int or not 1 <= limite <= 10:
        raise ValueError(
            "El limite debe ser un entero entre 1 y 10."
        )

    if criterio not in ("ingresos", "unidades"):
        raise ValueError(
            "El criterio debe ser ingresos o unidades."
        )

    detalle = obtener_detalle_productos()
    resumen = analizar_productos(detalle)

    columna = (
        "VentasProductos"
        if criterio == "ingresos"
        else "UnidadesVendidas"
    )

    top = (
        resumen.sort_values(
            [columna, "ProductID"],
            ascending=[False, True]
        )
        .head(limite)
    )

    productos = []

    for _, fila in top.iterrows():
        productos.append({
            "producto_id": int(fila["ProductID"]),
            "nombre": str(fila["NombreProducto"]),
            "unidades_vendidas": int(
                fila["UnidadesVendidas"]
            ),
            "cantidad_pedidos": int(
                fila["CantidadPedidos"]
            ),
            "ventas_productos": round(
                float(fila["VentasProductos"]), 2
            )
        })

    return {
        "criterio_ordenacion": criterio,
        "metrica_ingresos": "LineTotal",
        "moneda": "no especificada",
        "productos": productos
    }



def generar_reporte_ventas() -> dict:
    """Genera un reporte Excel con analisis empresarial."""

    print("\nConsultando datos para el reporte Excel...")

    # 1. Extraer datos desde SQL Server
    pedidos = obtener_pedidos()
    detalle = obtener_detalle_productos()

    # 2. Generar las cinco hojas originales
    ruta = generar_reporte_excel(
        pedidos=pedidos,
        detalle=detalle
    )

    if not ruta.is_file():
        raise RuntimeError(
            "El archivo Excel no se genero correctamente."
        )

    print(f"Reporte guardado en: {ruta}")

    # 3. Solicitar conclusiones a Gemini
    conclusiones_incluidas = False
    advertencia = None

    try:
        print("Generando conclusiones con Gemini...")

        conclusiones = generar_conclusiones_ia(
            pedidos,
            detalle
        )

        agregar_conclusiones_ia(
            ruta,
            conclusiones
        )

        conclusiones_incluidas = True

        print("Conclusiones incorporadas al Excel.")

    except Exception as error:
        # Mantener el reporte original si falla Gemini
        advertencia = (
            "El Excel fue generado, pero no se pudieron "
            "incorporar las conclusiones de IA. "
            f"Tipo de error: {type(error).__name__}"
        )

        print(advertencia)

    # 4. Preparar los metadatos del reporte
    hojas = [
        "Resumen",
        "Ventas mensuales",
        "Territorios",
        "Top ingresos",
        "Top unidades"
    ]

    if conclusiones_incluidas:
        hojas.append("Conclusiones IA")

    resultado = {
        "estado": "generado",
        "archivo": ruta.name,
        "carpeta": "reports",
        "formato": "xlsx",
        "hojas": hojas,
        "cantidad_hojas": len(hojas),
        "cantidad_graficos": 2,
        "conclusiones_ia_en_excel": conclusiones_incluidas
    }

    if advertencia:
        resultado["advertencia"] = advertencia

    return resultado


def generar_reporte_personalizado(
    anio: int | None = None,
    territorio: str | None = None,
    incluir_conclusiones_ia: bool = False
) -> dict:
    """
    Herramienta del agente para generar
    reportes personalizados.
    """

    # 1. Exigir al menos un filtro
    if anio is None and territorio is None:
        raise ValueError(
            "El reporte personalizado requiere "
            "un año o un territorio."
        )

    # 2. Validar tipos antes de consultar SQL
    if anio is not None:
        if (
            type(anio) is not int
            or not 1900 <= anio <= 2100
        ):
            raise ValueError(
                "El año debe ser un entero valido."
            )

    if territorio is not None:
        if (
            not isinstance(territorio, str)
            or not territorio.strip()
        ):
            raise ValueError(
                "El territorio no es valido."
            )

        territorio = territorio.strip()

    if type(incluir_conclusiones_ia) is not bool:
        raise ValueError(
            "incluir_conclusiones_ia debe "
            "ser verdadero o falso."
        )

    # 3. Generar el reporte filtrado
    ruta = (
        personalized_report.generar_reporte_personalizado(
            anio=anio,
            territorio=territorio,
            incluir_conclusiones_ia=incluir_conclusiones_ia
        )
    )

    if not ruta.is_file():
        raise RuntimeError(
            "El reporte no fue generado."
        )

    # 4. Verificar las hojas realmente guardadas
    wb = load_workbook(ruta, read_only=True)

    try:
        hojas = list(wb.sheetnames)
    finally:
        wb.close()

    tiene_conclusiones = (
        "Conclusiones IA" in hojas
    )

    # 5. Devolver informacion al agente
    resultado = {
        "estado": "generado",
        "archivo": ruta.name,
        "carpeta": "reports",
        "formato": "xlsx",
        "tipo_reporte": "personalizado",
        "filtros": {
            "anio": anio,
            "territorio": territorio
        },
        "hojas": hojas,
        "cantidad_hojas": len(hojas),
        "cantidad_graficos": 2,
        "conclusiones_ia_en_excel": tiene_conclusiones
    }

    if incluir_conclusiones_ia and not tiene_conclusiones:
        resultado["advertencia"] = (
            "El Excel fue generado, pero no "
            "se pudieron incorporar las "
            "conclusiones de IA."
        )

    return resultado