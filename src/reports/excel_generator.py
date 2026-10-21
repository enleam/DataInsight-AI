from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.analytics.sales_analysis import (
    preparar_pedidos,
    calcular_kpis,
    ventas_por_mes,
    ventas_por_territorio,
    analizar_productos
)


# Carpeta principal del proyecto
ROOT = Path(__file__).resolve().parents[2]

COLOR_AZUL = "17365D"
COLOR_BLANCO = "FFFFFF"
COLOR_VERDE = "008C76"

FORMATO_MONTO = "#,##0.00"
FORMATO_ENTERO = "#,##0"


def convertir_valor(valor):
    """Convierte tipos de Pandas a valores de Excel."""

    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()

    if isinstance(valor, Decimal):
        return float(valor)

    if hasattr(valor, "item"):
        return valor.item()

    return valor


def escribir_tabla(
    hoja,
    dataframe,
    encabezados,
    anchos,
    columnas_monetarias=()
):
    """Escribe y da formato a una tabla de datos."""

    if len(dataframe.columns) != len(encabezados):
        raise ValueError(
            "Las columnas no coinciden con los encabezados."
        )

    hoja.append(encabezados)

    # Agregar filas
    for fila in dataframe.itertuples(index=False, name=None):
        hoja.append([
            convertir_valor(valor)
            for valor in fila
        ])

    # Formato del encabezado
    for celda in hoja[1]:
        celda.font = Font(
            bold=True,
            color=COLOR_BLANCO
        )
        celda.fill = PatternFill(
            "solid",
            fgColor=COLOR_AZUL
        )
        celda.alignment = Alignment(
            horizontal="center"
        )

    # Ajustar anchos
    for indice, ancho in enumerate(anchos, start=1):
        letra = get_column_letter(indice)
        hoja.column_dimensions[letra].width = ancho

    # Formato numerico
    for fila in hoja.iter_rows(min_row=2):
        for indice, celda in enumerate(fila, start=1):

            if indice in columnas_monetarias:
                celda.number_format = FORMATO_MONTO

            elif isinstance(celda.value, int):
                celda.number_format = FORMATO_ENTERO

    # Opciones de visualizacion
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = (
        f"A1:{get_column_letter(len(encabezados))}"
        f"{hoja.max_row}"
    )

    hoja.sheet_view.showGridLines = False


def generar_reporte_excel(
    pedidos: pd.DataFrame,
    detalle: pd.DataFrame
) -> Path:
    """Genera un reporte empresarial en Excel."""

    # 1. Preparar datos
    pedidos = preparar_pedidos(pedidos)

    if pedidos.empty:
        raise ValueError("No existen pedidos para el reporte.")

    kpis = calcular_kpis(pedidos)
    mensuales = ventas_por_mes(pedidos)
    territorios = ventas_por_territorio(pedidos)
    productos = analizar_productos(detalle)

    top_ingresos = productos.head(10)

    top_unidades = (
        productos.sort_values(
            ["UnidadesVendidas", "ProductID"],
            ascending=[False, True]
        )
        .head(10)
    )

    # Comprobar consistencia monetaria
    diferencia = abs(
        float(kpis["total_ventas"])
        - float(detalle["LineTotal"].sum())
    )

    if diferencia > 0.01:
        raise ValueError(
            "Los importes de pedidos y detalles no coinciden."
        )

    # 2. Crear libro Excel
    wb = Workbook()

    # HOJA 1: RESUMEN EJECUTIVO
    resumen = wb.active
    resumen.title = "Resumen"

    resumen.merge_cells("A1:D2")
    titulo = resumen["A1"]
    titulo.value = "DATAINSIGHT AI - REPORTE DE VENTAS"
    titulo.font = Font(
        bold=True,
        size=16,
        color=COLOR_BLANCO
    )
    titulo.fill = PatternFill(
        "solid",
        fgColor=COLOR_AZUL
    )
    titulo.alignment = Alignment(
        vertical="center"
    )

    resumen.column_dimensions["A"].width = 30
    resumen.column_dimensions["B"].width = 68
    resumen.sheet_view.showGridLines = False

    indicadores = [
        (
            "Periodo",
            f"{kpis['primera_fecha'].date()} a "
            f"{kpis['ultima_fecha'].date()}"
        ),
        ("Total de pedidos", int(kpis["total_pedidos"])),
        ("Ventas de productos", float(kpis["total_ventas"])),
        ("Ticket promedio", float(kpis["ticket_promedio"]))
    ]

    for fila, (nombre, valor) in enumerate(
        indicadores,
        start=4
    ):
        resumen.cell(fila, 1, nombre)
        resumen.cell(fila, 2, valor)

        resumen.cell(fila, 1).font = Font(
            bold=True,
            color=COLOR_AZUL
        )

    resumen["B5"].number_format = FORMATO_ENTERO
    resumen["B6"].number_format = FORMATO_MONTO
    resumen["B7"].number_format = FORMATO_MONTO

    resumen["A10"] = "HALLAZGOS PRINCIPALES"
    resumen["A10"].font = Font(
        bold=True,
        color=COLOR_VERDE,
        size=12
    )

    lider_territorio = territorios.iloc[0]
    lider_producto = top_ingresos.iloc[0]

    resumen["A12"] = "Territorio lider"
    resumen["B12"] = str(lider_territorio["Territorio"])

    resumen["A13"] = "Producto lider en ingresos"
    resumen["B13"] = str(lider_producto["NombreProducto"])

    resumen["A15"] = "Nota metodologica"
    resumen["B15"] = (
        "Ventas de productos sin impuestos ni transporte. "
        "Los años 2011 y 2014 contienen periodos parciales."
    )

    resumen["A16"] = "Moneda"
    resumen["B16"] = "No especificada en el reporte"

    for fila in [12, 13, 15, 16]:
        resumen.cell(fila, 1).font = Font(bold=True)
        resumen.cell(fila, 2).alignment = Alignment(
            wrap_text=True,
            vertical="center"
        )

    resumen.row_dimensions[15].height = 32

    # HOJA 2: VENTAS MENSUALES
    hoja_meses = wb.create_sheet("Ventas mensuales")

    escribir_tabla(
        hoja_meses,
        mensuales,
        ["Mes", "Pedidos", "Ventas de productos"],
        [16, 18, 24],
        columnas_monetarias=(3,)
    )

    grafico_mensual = LineChart()
    grafico_mensual.title = "Evolucion mensual de ventas"
    grafico_mensual.y_axis.title = "Ingresos"
    grafico_mensual.x_axis.title = "Mes"
    grafico_mensual.width = 19
    grafico_mensual.height = 10
    grafico_mensual.legend = None

    grafico_mensual.add_data(
        Reference(
            hoja_meses,
            min_col=3,
            min_row=1,
            max_row=hoja_meses.max_row
        ),
        titles_from_data=True
    )

    grafico_mensual.set_categories(
        Reference(
            hoja_meses,
            min_col=1,
            min_row=2,
            max_row=hoja_meses.max_row
        )
    )

    hoja_meses.add_chart(grafico_mensual, "E3")

    # HOJA 3: TERRITORIOS
    hoja_territorios = wb.create_sheet("Territorios")

    escribir_tabla(
        hoja_territorios,
        territorios,
        ["Territorio", "Pedidos", "Ventas de productos"],
        [24, 18, 25],
        columnas_monetarias=(3,)
    )

    grafico_territorios = BarChart()
    grafico_territorios.type = "bar"
    grafico_territorios.title = "Ingresos por territorio"
    grafico_territorios.width = 19
    grafico_territorios.height = 11
    grafico_territorios.legend = None

    grafico_territorios.add_data(
        Reference(
            hoja_territorios,
            min_col=3,
            min_row=1,
            max_row=hoja_territorios.max_row
        ),
        titles_from_data=True
    )

    grafico_territorios.set_categories(
        Reference(
            hoja_territorios,
            min_col=1,
            min_row=2,
            max_row=hoja_territorios.max_row
        )
    )

    hoja_territorios.add_chart(
        grafico_territorios,
        "E3"
    )

    # HOJA 4: TOP PRODUCTOS POR INGRESOS
    columnas_productos = [
        "ProductID",
        "NombreProducto",
        "UnidadesVendidas",
        "CantidadPedidos",
        "VentasProductos"
    ]

    encabezados = [
        "ID",
        "Producto",
        "Unidades vendidas",
        "Pedidos",
        "Ingresos"
    ]

    anchos = [12, 38, 22, 18, 25]

    hoja_ingresos = wb.create_sheet("Top ingresos")

    escribir_tabla(
        hoja_ingresos,
        top_ingresos[columnas_productos],
        encabezados,
        anchos,
        columnas_monetarias=(5,)
    )

    # HOJA 5: TOP PRODUCTOS POR UNIDADES
    hoja_unidades = wb.create_sheet("Top unidades")

    escribir_tabla(
        hoja_unidades,
        top_unidades[columnas_productos],
        encabezados,
        anchos,
        columnas_monetarias=(5,)
    )

    # 3. Guardar el reporte
    carpeta = ROOT / "reports"
    carpeta.mkdir(parents=True, exist_ok=True)

    fecha = datetime.now().strftime("%Y%m%d_%H%M%S")

    ruta = carpeta / f"DataInsightAI_ventas_{fecha}.xlsx"

    wb.save(ruta)

    return ruta

def agregar_conclusiones_ia(
    ruta: Path,
    conclusiones: list[str]
) -> None:
    """Agrega conclusiones generadas por IA al Excel."""

    if (
        len(conclusiones) != 3
        or any(
            not isinstance(texto, str) or not texto.strip()
            for texto in conclusiones
        )
    ):
        raise ValueError(
            "Se requieren exactamente tres conclusiones."
        )

    # 1. Abrir el archivo ya generado
    wb = load_workbook(ruta)

    # 2. Crear la nueva hoja
    hoja = wb.create_sheet("Conclusiones IA")

    hoja.sheet_view.showGridLines = False

    hoja.column_dimensions["A"].width = 90
    hoja.column_dimensions["B"].width = 15
    hoja.column_dimensions["C"].width = 15
    hoja.column_dimensions["D"].width = 15

    # 3. Titulo
    hoja.merge_cells("A1:D2")

    titulo = hoja["A1"]
    titulo.value = "DATAINSIGHT AI - CONCLUSIONES"

    titulo.font = Font(
        size=16,
        bold=True,
        color=COLOR_BLANCO
    )

    titulo.fill = PatternFill(
        "solid",
        fgColor=COLOR_AZUL
    )

    titulo.alignment = Alignment(
        vertical="center"
    )

    # 4. Incorporar conclusiones
    for numero, conclusion in enumerate(
        conclusiones,
        start=1
    ):
        fila = numero + 3

        hoja.merge_cells(
            start_row=fila,
            start_column=1,
            end_row=fila,
            end_column=4
        )

        celda = hoja.cell(
            row=fila,
            column=1
        )

        celda.value = f"{numero}. {conclusion}"
        celda.data_type = "s"

        celda.alignment = Alignment(
            wrap_text=True,
            vertical="center"
        )

        hoja.row_dimensions[fila].height = 65

    # 5. Advertencia metodologica
    hoja["A9"] = "Nota metodologica"
    hoja["A9"].font = Font(
        bold=True,
        color=COLOR_VERDE
    )

    hoja.merge_cells("A10:D11")

    hoja["A10"] = (
        "Estas conclusiones fueron redactadas mediante "
        "inteligencia artificial a partir de indicadores "
        "calculados con Python y Pandas. "
        "Deben revisarse antes de tomar decisiones."
    )

    hoja["A10"].alignment = Alignment(
        wrap_text=True,
        vertical="center"
    )

    # 6. Guardar el archivo actualizado
    wb.save(ruta)